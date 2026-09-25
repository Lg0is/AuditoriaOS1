import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

arquivo_base = 'Base_Unificada.xlsx'
arquivo_auditoria = 'Att.02 - Auditoria de ordem de serviço - SLG.xlsx'
arquivo_saida = 'Relatorio_Auditoria_OS.xlsx'

coluna_os_df = 'Ordem de Serviço'
coluna_os_df2 = 'Ordem de Serviço'

df = pd.read_excel(arquivo_base)
df2 = pd.read_excel(arquivo_auditoria)

df.columns = df.columns.str.strip()
df2.columns = df2.columns.str.strip()


def padronizar_os(serie):
    return (
        serie
        .astype('string')
        .str.strip()
        .str.replace(r'\.0$', '', regex=True)
    )


df[coluna_os_df] = padronizar_os(df[coluna_os_df])
df2[coluna_os_df2] = padronizar_os(df2[coluna_os_df2])

df['Status da Atividade'] = (
    df['Status da Atividade']
    .astype('string')
    .str.strip()
    .str.upper()
)

status_validos = [
    'CONCLUÍDO',
    'NÃO CONCLUÍDO'
]

df_servicos = df.loc[
    df['Status da Atividade'].isin(status_validos),
    [
        'Recurso',
        coluna_os_df,
        'Status da Atividade'
    ]
].copy()

df_servicos = df_servicos.dropna(
    subset=[
        'Recurso',
        coluna_os_df
    ]
)

df_servicos = df_servicos[
    df_servicos[coluna_os_df].ne('')
]

df_servicos = df_servicos.drop_duplicates(
    subset=[
        'Recurso',
        coluna_os_df
    ],
    keep='last'
)

campos_auditoria = [
    coluna_os_df2,
    'Protocolo ok?',
    'Foto da fachada ok?',
    'A foto comprova resultado na O.S.?',
    'Parecer de execução explicativo ok?',
    'HD cadastrado no Sansys igual foto?',
    'Motivo de não execução coerente?',
    'Há necessidade ação corretiva posterior?',
    'Qual ação?'
]

campos_nao_encontrados = [
    coluna
    for coluna in campos_auditoria
    if coluna not in df2.columns
]

if campos_nao_encontrados:
    print('\nCOLUNAS DISPONÍVEIS NA SEGUNDA BASE:\n')
    print('\n'.join(df2.columns.astype(str)))

    raise KeyError(
        '\nColunas não encontradas na segunda base: '
        + ', '.join(campos_nao_encontrados)
    )

df2_selecionado = df2[
    campos_auditoria
].copy()

df2_selecionado = df2_selecionado.dropna(
    subset=[coluna_os_df2]
)

df2_selecionado = df2_selecionado[
    df2_selecionado[coluna_os_df2].ne('')
]

df2_selecionado = df2_selecionado.drop_duplicates(
    subset=[coluna_os_df2],
    keep='last'
)

relatorio = df_servicos.merge(
    df2_selecionado,
    left_on=coluna_os_df,
    right_on=coluna_os_df2,
    how='inner'
)

if coluna_os_df != coluna_os_df2:
    relatorio = relatorio.drop(
        columns=[coluna_os_df2]
    )

relatorio = relatorio.rename(
    columns={
        'Recurso': 'Equipe',
        coluna_os_df: 'Ordem de Serviço',
        'Status da Atividade': 'Status'
    }
)

ordem_colunas = [
    'Equipe',
    'Ordem de Serviço',
    'Status',
    'Protocolo ok?',
    'Foto da fachada ok?',
    'A foto comprova resultado na O.S.?',
    'Parecer de execução explicativo ok?',
    'HD cadastrado no Sansys igual foto?',
    'Motivo de não execução coerente?',
    'Há necessidade ação corretiva posterior?',
    'Qual ação?'
]

relatorio = (
    relatorio[ordem_colunas]
    .sort_values(
        by=[
            'Equipe',
            'Status',
            'Ordem de Serviço'
        ],
        ascending=[
            True,
            True,
            True
        ]
    )
    .reset_index(drop=True)
)

resumo_por_equipe = (
    relatorio
    .groupby('Equipe')
    .agg(
        Concluidos=(
            'Status',
            lambda valores: valores.eq(
                'CONCLUÍDO'
            ).sum()
        ),
        Nao_Concluidos=(
            'Status',
            lambda valores: valores.eq(
                'NÃO CONCLUÍDO'
            ).sum()
        ),
        Acao_Corretiva_Necessaria=(
            'Há necessidade ação corretiva posterior?',
            lambda valores: (
                valores
                .astype('string')
                .str.strip()
                .str.upper()
                .eq('SIM')
                .sum()
            )
        )
    )
    .reset_index()
)

resumo_por_equipe['Total'] = (
    resumo_por_equipe['Concluidos']
    + resumo_por_equipe['Nao_Concluidos']
)

resumo_por_equipe = (
    resumo_por_equipe[
        [
            'Equipe',
            'Concluidos',
            'Nao_Concluidos',
            'Total',
            'Acao_Corretiva_Necessaria'
        ]
    ]
    .sort_values(
        by='Total',
        ascending=False
    )
    .reset_index(drop=True)
)

total_concluidos = resumo_por_equipe[
    'Concluidos'
].sum()

total_nao_concluidos = resumo_por_equipe[
    'Nao_Concluidos'
].sum()

total_servicos = resumo_por_equipe[
    'Total'
].sum()

total_acoes_corretivas = resumo_por_equipe[
    'Acao_Corretiva_Necessaria'
].sum()

linha_total = pd.DataFrame(
    {
        'Equipe': ['TOTAL GERAL'],
        'Concluidos': [total_concluidos],
        'Nao_Concluidos': [total_nao_concluidos],
        'Total': [total_servicos],
        'Acao_Corretiva_Necessaria': [
            total_acoes_corretivas
        ]
    }
)

resumo_final = pd.concat(
    [
        resumo_por_equipe,
        linha_total
    ],
    ignore_index=True
)

os_nao_encontradas = df_servicos.merge(
    df2_selecionado[[coluna_os_df2]],
    left_on=coluna_os_df,
    right_on=coluna_os_df2,
    how='left',
    indicator=True
)

os_nao_encontradas = (
    os_nao_encontradas[
        os_nao_encontradas['_merge'].eq(
            'left_only'
        )
    ][
        [
            'Recurso',
            coluna_os_df,
            'Status da Atividade'
        ]
    ]
    .rename(
        columns={
            'Recurso': 'Equipe',
            coluna_os_df: 'Ordem de Serviço',
            'Status da Atividade': 'Status'
        }
    )
    .drop_duplicates()
    .sort_values(
        by=[
            'Equipe',
            'Status',
            'Ordem de Serviço'
        ]
    )
    .reset_index(drop=True)
)

with pd.ExcelWriter(
    arquivo_saida,
    engine='openpyxl'
) as writer:

    relatorio.to_excel(
        writer,
        sheet_name='Auditoria por OS',
        index=False
    )

    resumo_final.to_excel(
        writer,
        sheet_name='Resumo por Equipe',
        index=False
    )

    os_nao_encontradas.to_excel(
        writer,
        sheet_name='OS Não Encontradas',
        index=False
    )

    for planilha in writer.sheets.values():

        planilha.freeze_panes = 'A2'
        planilha.auto_filter.ref = planilha.dimensions
        planilha.sheet_view.showGridLines = False
        planilha.row_dimensions[1].height = 45

        for celula in planilha[1]:

            celula.font = Font(
                bold=True,
                color='FFFFFF'
            )

            celula.fill = PatternFill(
                fill_type='solid',
                fgColor='1F4E78'
            )

            celula.alignment = Alignment(
                horizontal='center',
                vertical='center',
                wrap_text=True
            )

        for linha in planilha.iter_rows(
            min_row=2
        ):
            for celula in linha:
                celula.alignment = Alignment(
                    vertical='top',
                    wrap_text=True
                )

        for numero_coluna, coluna in enumerate(
            planilha.columns,
            start=1
        ):

            maior_tamanho = max(
                (
                    len(str(celula.value))
                    for celula in coluna
                    if celula.value is not None
                ),
                default=0
            )

            letra = get_column_letter(
                numero_coluna
            )

            planilha.column_dimensions[
                letra
            ].width = min(
                max(
                    maior_tamanho + 2,
                    14
                ),
                40
            )

    planilha_auditoria = writer.sheets[
        'Auditoria por OS'
    ]

    larguras_auditoria = {
        'A': 25,
        'B': 20,
        'C': 20,
        'D': 18,
        'E': 25,
        'F': 32,
        'G': 35,
        'H': 32,
        'I': 35,
        'J': 38,
        'K': 45
    }

    for coluna, largura in larguras_auditoria.items():
        planilha_auditoria.column_dimensions[
            coluna
        ].width = largura

    planilha_resumo = writer.sheets[
        'Resumo por Equipe'
    ]

    ultima_linha_resumo = (
        planilha_resumo.max_row
    )

    for celula in planilha_resumo[
        ultima_linha_resumo
    ]:

        celula.font = Font(
            bold=True
        )

        celula.fill = PatternFill(
            fill_type='solid',
            fgColor='D9EAD3'
        )

        celula.alignment = Alignment(
            horizontal='center',
            vertical='center'
        )

print(f'\nRELATÓRIO GERADO: {arquivo_saida}')

print('\nAUDITORIA ORGANIZADA POR EQUIPE E OS:\n')
print(
    relatorio.to_string(
        index=False
    )
)

print('\nRESUMO POR EQUIPE:\n')
print(
    resumo_final.to_string(
        index=False
    )
)

print(
    f'\nTOTAL DE SERVIÇOS CONCLUÍDOS: '
    f'{total_concluidos}'
)

print(
    f'TOTAL DE SERVIÇOS NÃO CONCLUÍDOS: '
    f'{total_nao_concluidos}'
)

print(
    f'TOTAL GERAL DE SERVIÇOS: '
    f'{total_servicos}'
)

print(
    f'TOTAL DE OS ENCONTRADAS NA AUDITORIA: '
    f'{relatorio["Ordem de Serviço"].nunique()}'
)

print(
    f'TOTAL DE OS NÃO ENCONTRADAS: '
    f'{os_nao_encontradas["Ordem de Serviço"].nunique()}'
)